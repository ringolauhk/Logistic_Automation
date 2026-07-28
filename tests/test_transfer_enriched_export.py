"""Final enriched product lines: shared schema, UI wiring, and the Excel
export (Build 10). Fully offline - the export module has no transport,
no auth, and touches no workflow state."""

import json
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook

from apps.web.job_manager import JobError
from apps.web.transfer import enriched_export as ex
from apps.web.transfer import jobs as tjobs
from apps.web.transfer import product_lookup as pl
from tests.test_transfer_product_lookup import (
    EAN_A,
    EAN_B,
    approved_job,
    envelope,
    run_lookup,
    wire_record,
)

ROOT = Path(__file__).resolve().parent.parent


@pytest.fixture(autouse=True)
def roots(tmp_path, monkeypatch):
    monkeypatch.setenv("WEB_JOBS_DIR", str(tmp_path / "jobs"))
    monkeypatch.setenv("TRANSFER_JOBS_DIR", str(tmp_path / "transfer-jobs"))
    return tmp_path


RECORD_A = wire_record(EAN_A, **{
    **{f"analysisCode{i:02d}": "" for i in range(1, 16)},
    "analysisCode01": "FAKE SEASON", "analysisCode06": "12.5",
    **{f"compositon{i}": "" for i in range(1, 5)},
    "compositon1": "FAKE 90% WOOL"})
RECORD_B = wire_record(EAN_B, item="ZETF381237E085", color="E085",
                       size="XS", desc="SRT - JAZZ SHORTS", price=1900.00)


def enriched_job(tmp_path):
    job_id = approved_job(tmp_path)
    run_lookup(job_id, [(200, envelope([RECORD_A, RECORD_B]))])
    return job_id


# --- shared schema ----------------------------------------------------------------

class TestSchema:
    def test_stable_full_schema(self):
        headers = [h for h, _ in ex.EXPORT_COLUMNS]
        for required in ("Line", "Carton", "Status", "Via", "Issues",
                         "Source EAN", "Source item", "Source description",
                         "Source quantity", "API EAN", "API item",
                         "API color", "API size", "API description",
                         "Original price", "Discount price"):
            assert required in headers, required
        assert [h for h in headers if h.startswith("AC")] == \
            [f"AC{i:02d}" for i in range(1, 16)]
        assert [h for h in headers if h.startswith("Composition")] == \
            [f"Composition {i}" for i in range(1, 5)]
        # prices come BEFORE the attribute block; identity stays left
        assert headers.index("Original price") < headers.index("AC01")
        assert headers.index("Line") < headers.index("Source EAN") < \
            headers.index("AC01")
        # the misspelled wire name never reaches the UI schema
        assert not any("compositon" in h.lower() for h in headers)

    def test_rows_one_per_line_with_attributes_inline(self, tmp_path):
        job_id = enriched_job(tmp_path)
        enrichment = pl.load_enrichment(job_id)
        rows = ex.build_rows(enrichment)
        assert len(rows) == len(enrichment["line_enrichments"]) == 2
        by_ean = {r["Source EAN"]: r for r in rows}
        row_a = by_ean[EAN_A]
        assert row_a["AC01"] == "FAKE SEASON"
        assert row_a["AC06"] == "12.5"
        assert row_a["AC02"] == ""                       # blank stays blank
        assert row_a["Composition 1"] == "FAKE 90% WOOL"
        assert row_a["Original price"] == 1400.0         # numeric
        assert isinstance(row_a["Source quantity"], int)
        assert row_a["API EAN"] == EAN_A                 # leading zero str
        assert isinstance(row_a["API EAN"], str)

    def test_duplicate_products_on_distinct_lines_stay_separate(
            self, tmp_path):
        from tests.test_transfer_extraction import ROW_A
        dup_rows = (ROW_A, ("2",) + ROW_A[1:])
        job_id = approved_job(tmp_path,
                              [{"rows": dup_rows, "carton_total": None}])
        run_lookup(job_id, [(200, envelope([RECORD_A]))])
        rows = ex.build_rows(pl.load_enrichment(job_id))
        assert len(rows) == 2                            # one per line
        assert rows[0]["Line"] != rows[1]["Line"]
        assert rows[0]["API EAN"] == rows[1]["API EAN"]


# --- Excel export -----------------------------------------------------------------

class TestExcelExport:
    def test_valid_workbook_structure_and_formats(self, tmp_path):
        job_id = enriched_job(tmp_path)
        data = ex.build_enriched_workbook_bytes(job_id)
        book = load_workbook(BytesIO(data))              # opens cleanly
        sheet = book[ex.WORKSHEET_NAME]
        assert book.sheetnames == [ex.WORKSHEET_NAME]
        assert sheet.freeze_panes == "A2"
        assert sheet.auto_filter.ref is not None
        headers = [c.value for c in sheet[1]]
        assert headers == [h for h, _ in ex.EXPORT_COLUMNS]
        assert all(c.font.bold for c in sheet[1])
        assert sheet.max_row == 3                        # header + 2 lines
        ean_col = headers.index("Source EAN") + 1
        cell = sheet.cell(row=2, column=ean_col)
        assert isinstance(cell.value, str)
        assert cell.value.startswith("0")                # leading zero kept
        assert cell.number_format == "@"
        qty_col = headers.index("Source quantity") + 1
        assert isinstance(sheet.cell(row=2, column=qty_col).value, int)
        price_col = headers.index("Original price") + 1
        price_value = sheet.cell(row=2, column=price_col).value
        assert isinstance(price_value, (int, float))     # numeric, not text
        assert not isinstance(price_value, str)
        assert float(price_value) == 1400.0
        ac2_col = headers.index("AC02") + 1
        assert sheet.cell(row=2, column=ac2_col).value is None   # blank cell
        assert ex.export_filename(job_id) == \
            f"Transfer_{job_id}_Enriched_Product_Lines.xlsx"

    def test_export_row_count_matches_enrichment(self, tmp_path):
        job_id = enriched_job(tmp_path)
        enrichment = pl.load_enrichment(job_id)
        book = load_workbook(
            BytesIO(ex.build_enriched_workbook_bytes(job_id)))
        assert (book[ex.WORKSHEET_NAME].max_row - 1
                == len(enrichment["line_enrichments"]))

    def test_export_changes_no_state_and_no_checkpoints(self, tmp_path):
        job_id = enriched_job(tmp_path)
        status_before = tjobs.load_transfer_job(job_id).status
        artifact_before = pl.result_path(job_id).read_bytes()
        ex.build_enriched_workbook_bytes(job_id)
        assert tjobs.load_transfer_job(job_id).status == status_before
        assert pl.result_path(job_id).read_bytes() == artifact_before
        # nothing written into the job folder by the export
        job_dir = tjobs.transfer_job_dir_for(job_id)
        assert not list(job_dir.rglob("*.xlsx"))
        assert not list(job_dir.rglob("*.tmp-*"))

    def test_export_makes_no_api_or_auth_calls(self):
        src = (ROOT / "apps" / "web" / "transfer"
               / "enriched_export.py").read_text(encoding="utf-8")
        for forbidden in ("httpx", "build_client", "transport",
                          "ensure_access_token", "Authorization",
                          "run_product_lookup", "update_job_status"):
            assert forbidden not in src, forbidden

    def test_missing_enrichment_rejected(self, tmp_path):
        job_id = approved_job(tmp_path)
        with pytest.raises(JobError, match="No product lookup result"):
            ex.build_enriched_workbook_bytes(job_id)


# --- UI wiring --------------------------------------------------------------------

class TestUiWiring:
    RPAGE = (ROOT / "apps" / "web" / "transfer"
             / "review_page.py").read_text(encoding="utf-8")

    def test_single_final_table_replaces_attribute_tables(self):
        assert self.RPAGE.count("Final enriched product lines") == 1
        assert "All populated Analysis Codes" not in self.RPAGE
        assert "Per-line enrichment (source vs API)" not in self.RPAGE

    def test_column_selector_and_show_all_toggle(self):
        assert "Show all product attributes" in self.RPAGE
        assert "transfer_enriched_columns" in self.RPAGE
        assert "DEFAULT_VISIBLE_COLUMNS" in self.RPAGE
        # one table only: a single dataframe call in the enriched renderer
        renderer = self.RPAGE.split("_render_enriched_lines_table")[2]
        renderer = renderer.split("def ")[0]
        assert renderer.count("st.dataframe") == 1

    def test_excel_download_wired_with_correct_name(self):
        assert "Download enriched product lines - Excel" in self.RPAGE
        assert "build_enriched_workbook_bytes" in self.RPAGE
        assert "export_filename" in self.RPAGE
        assert "spreadsheetml" in self.RPAGE
        # no CSV export for this table
        renderer = self.RPAGE.split("def _render_enriched_lines_table")[1]
        renderer = renderer.split("\ndef ")[0]
        assert "csv" not in renderer.lower()
