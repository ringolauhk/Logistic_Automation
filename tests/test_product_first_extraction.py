"""M11: product-first extraction for general product documents.

The scanner is not limited to payable invoices - users upload packing lists,
delivery notes and free-goods support lists whose payload is the PRODUCT
ROWS. The confirmed live failure: a "Free Goods Support Item List" carrying
seven complete product lines was failed outright because it printed no
seller name, invoice number or invoice date.

Policy under test:
  HARD - at least one usable product row (item code, barcode or description).
  SOFT - all document metadata and most per-row fields: warnings, exported
         blank and highlighted, never a failure and never an extra call.
Semantic safety (M10) stays hard: a MISSING value warns, a CONTRADICTORY
value still rejects the attempt.

Synthetic fixtures only - fake vendors, codes and amounts.
"""

import json
from pathlib import Path

import pytest

from invoice_extractor import openrouter_client
from invoice_extractor.excel_export import (
    INVOICE_COLUMNS,
    LINE_ITEM_COLUMNS,
    export_workbook,
)
from invoice_extractor.pipeline import process_file, safe_review_categories
from invoice_extractor.schema import (
    ExtractionError,
    check_extractable,
    declares_free_of_charge,
    detect_document_type,
    document_missing_fields,
    line_item_is_usable,
    line_item_missing_fields,
    normalize_invoice,
    usable_line_items,
)

from .conftest import build_pdf, invoice_dict, make_config
from .test_openrouter_pipeline import Recorder, envelope
from .test_vision_fallback import text_calls, vision_calls

# A free-goods support list: products priced but declared free of charge,
# no seller / invoice number / invoice date. Fake values throughout.
FREE_GOODS_ROWS = [
    ("50000001", "AL NIGHT CREME DLX 15ML", 200, 30, 6000),
    ("50000002", "AL SHAMPOO SAMPLE 7ML X2", 700, 14, 9800),
    ("50000003", "AL HAIR OIL SAMPLE 5ML", 700, 7, 4900),
]
FG_QTY_TOTAL = 1600
FG_AMOUNT_TOTAL = 20700

FREE_GOODS_TEXT = (
    "Free Goods Support Item List\n"
    "The item(s) listed above is/are free-of-charge\n"
    "S/O NO.: 9000001\nDELIVERY DATE: 06.08.2026\n"
    "SOLD-TO: 10000001\nEXAMPLE BUYER LIMITED\n"
    "MATERIAL DESCRIPTION QUANTITY UNIT PRICE AMOUNT (HK$)\n"
    + "\n".join(f"{c} {d} {q} {u}.00 {a}.00"
                for c, d, q, u, a in FREE_GOODS_ROWS)
    + "\n- TOTAL - 20,700.00\n"
)
PACKING_LIST_TEXT = (
    "Packing List\nDELIVERY NOTE NO 7788\nCarton 1 of 3\n"
    "ITEM CODE DESCRIPTION QUANTITY\n"
    "PK-1 Widget blue 10\nPK-2 Widget red 5\n"
)


def product_doc(**overrides) -> dict:
    """A product document with NO seller/date/invoice number - only rows."""
    data = invoice_dict(
        invoice_number=None, invoice_date=None, seller_name=None,
        seller_address=None, buyer_name="Example Buyer Limited",
        currency="HKD", subtotal=None, tax_amount=None,
        total_amount=FG_AMOUNT_TOTAL, payment_terms=None,
        line_items=[{"item_code": c, "description": d, "quantity": q,
                     "unit_price": u, "amount": a}
                    for c, d, q, u, a in FREE_GOODS_ROWS])
    data.update(overrides)
    return data


def product_json(**overrides) -> str:
    return json.dumps(product_doc(**overrides))


def pf_cfg(n_text_models=2, **overrides):
    base = dict(
        llm_gateway="openrouter",
        openrouter_api_key="test-or-key",
        openrouter_text_models=tuple(
            f"test-vendor/text-{i + 1}" for i in range(n_text_models)),
        openrouter_vision_models=("test-vendor/vision-1",),
        max_retries=1,
    )
    base.update(overrides)
    return make_config(**base)


@pytest.fixture
def free_goods_pdf(tmp_path):
    return Path(build_pdf(tmp_path / "free_goods.pdf",
                          [("text", FREE_GOODS_TEXT)]))


def _inv(**overrides):
    return normalize_invoice(product_doc(**overrides))


# --- the hard/soft boundary ---------------------------------------------------

class TestUsabilityGate:
    @pytest.mark.parametrize("row,usable", [
        ({"item_code": "A1"}, True),                     # code only
        ({"barcode": "4900000000017"}, True),            # barcode only
        ({"description": "Widget"}, True),               # description only
        ({"quantity": 5, "unit_price": 2, "amount": 10}, False),   # no identity
        ({"item_code": "  ", "description": ""}, False),           # blank
        ({}, False),                                                # empty
    ])
    def test_row_usability(self, row, usable):
        inv = normalize_invoice({"line_items": [row]})
        if not inv.line_items:                    # fully empty row is dropped
            assert usable is False
            return
        assert line_item_is_usable(inv.line_items[0]) is usable

    def test_usable_rows_accepted_without_any_metadata(self):
        inv = _inv()
        check_extractable(inv)                    # no raise
        assert len(usable_line_items(inv)) == 3

    def test_no_usable_rows_fails(self):
        inv = normalize_invoice({"total_amount": 100, "line_items": [
            {"quantity": 5, "unit_price": 2, "amount": 10}]})
        with pytest.raises(ExtractionError, match="no usable product rows"):
            check_extractable(inv)

    def test_no_rows_at_all_fails(self):
        with pytest.raises(ExtractionError, match="no product rows"):
            check_extractable(normalize_invoice({"total_amount": 100}))

    def test_missing_document_fields_listed_not_raised(self):
        missing = document_missing_fields(_inv())
        for expected in ("invoice_number", "invoice_date", "seller_name"):
            assert expected in missing
        assert "currency" not in missing          # present in this fixture

    def test_missing_line_fields_listed_per_row(self):
        inv = _inv(line_items=[
            {"item_code": "A1", "description": "X", "quantity": None,
             "unit_price": 2, "amount": None}])
        missing = line_item_missing_fields(inv.line_items[0])
        assert "quantity" in missing and "amount" in missing
        assert "barcode" in missing
        assert "item_code" not in missing


# --- document type + free-of-charge evidence ---------------------------------

class TestDocumentEvidence:
    @pytest.mark.parametrize("text,expected", [
        (FREE_GOODS_TEXT, "free_goods_support_list"),
        (PACKING_LIST_TEXT, "packing_list"),
        ("Delivery Note 123\nItems below", "delivery_note"),
        ("Commercial Invoice\nSeller: X", "commercial_invoice"),
        ("TAX INVOICE 99", "invoice"),
        ("Statement of account", None),           # no explicit evidence
        (None, None),
    ])
    def test_document_type_from_explicit_wording(self, text, expected):
        assert detect_document_type(text) == expected

    @pytest.mark.parametrize("text,expected", [
        (FREE_GOODS_TEXT, True),
        ("Samples - no commercial value", True),
        ("Goods sold as per terms", False),
    ])
    def test_free_of_charge_evidence(self, text, expected):
        assert declares_free_of_charge(text) is expected


# --- pipeline behavior --------------------------------------------------------

class TestProductFirstPipeline:
    def test_free_goods_list_extracted_with_warnings(self, logger,
                                                     free_goods_pdf,
                                                     monkeypatch):
        rec = Recorder([envelope(product_json())])
        monkeypatch.setattr(openrouter_client, "_chat_completion", rec)

        result = process_file(free_goods_pdf, pf_cfg(), logger)

        assert len(rec.calls) == 1                # accepted first time
        assert vision_calls(rec) == []            # no fallback for metadata
        assert result.error is False              # EXTRACTED, not failed
        assert result.needs_review is True        # with warnings
        assert len(result.invoice.line_items) == 3
        assert sum(i.quantity for i in result.invoice.line_items) == FG_QTY_TOTAL
        assert sum(i.amount for i in result.invoice.line_items) == FG_AMOUNT_TOTAL
        assert result.invoice.currency == "HKD"   # HK$ evidence in the text
        assert result.invoice.seller_name is None      # blank, not invented
        assert result.invoice.invoice_date is None
        assert result.invoice.invoice_number is None
        assert result.document_type == "free_goods_support_list"
        assert result.free_of_charge is True
        for field in ("seller_name", "invoice_date", "invoice_number"):
            assert field in result.document_missing
        cats = safe_review_categories(result)
        assert "missing_document_metadata" in cats
        assert "provider_failure" not in cats

    @pytest.mark.parametrize("body,doc_type", [
        (PACKING_LIST_TEXT, "packing_list"),
        ("Delivery Note 55\nITEM QTY\nD-1 Widget 3", "delivery_note"),
        ("Goods list\nITEM QTY\nX-1 Widget 3", None),   # unknown type
    ])
    def test_other_product_documents_succeed(self, logger, tmp_path,
                                             monkeypatch, body, doc_type):
        pdf = Path(build_pdf(tmp_path / "doc.pdf", [("text", body)]))
        rec = Recorder([envelope(product_json(
            currency=None, total_amount=None))])
        monkeypatch.setattr(openrouter_client, "_chat_completion", rec)

        result = process_file(pdf, pf_cfg(), logger)

        assert len(rec.calls) == 1
        assert result.error is False
        assert result.needs_review is True
        assert len(result.invoice.line_items) == 3
        assert result.document_type == doc_type

    @pytest.mark.parametrize("absent", [
        "seller_name", "invoice_date", "currency", "total_amount"])
    def test_single_missing_header_never_rejects(self, logger, free_goods_pdf,
                                                 monkeypatch, absent):
        rec = Recorder([envelope(product_json(**{absent: None}))])
        monkeypatch.setattr(openrouter_client, "_chat_completion", rec)

        result = process_file(free_goods_pdf, pf_cfg(), logger)

        assert len(rec.calls) == 1                # NO escalation
        assert result.error is False
        assert len(result.invoice.line_items) == 3
        assert absent in result.document_missing

    def test_many_missing_fields_produce_one_warning_set(self, logger,
                                                         free_goods_pdf,
                                                         monkeypatch):
        rec = Recorder([envelope(product_json(
            currency=None, total_amount=None, buyer_name=None))])
        monkeypatch.setattr(openrouter_client, "_chat_completion", rec)

        result = process_file(free_goods_pdf, pf_cfg(), logger)

        assert result.review_reason.count("missing document metadata") == 1
        assert {"seller_name", "invoice_date", "currency", "total_amount",
                "buyer_name"} <= set(result.document_missing)

    @pytest.mark.parametrize("absent", ["quantity", "unit_price", "amount"])
    def test_missing_line_field_exports_row_with_warning(self, logger,
                                                         free_goods_pdf,
                                                         monkeypatch, absent):
        rows = [dict(zip(("item_code", "description", "quantity",
                          "unit_price", "amount"), r))
                for r in FREE_GOODS_ROWS]
        rows[0][absent] = None
        rec = Recorder([envelope(product_json(line_items=rows))])
        monkeypatch.setattr(openrouter_client, "_chat_completion", rec)

        result = process_file(free_goods_pdf, pf_cfg(), logger)

        assert len(rec.calls) == 1                # no extra provider call
        assert result.error is False
        assert len(result.invoice.line_items) == 3          # row kept
        assert getattr(result.invoice.line_items[0], absent) is None
        assert absent in result.line_missing[1]

    def test_document_with_no_usable_rows_fails(self, logger, free_goods_pdf,
                                                monkeypatch):
        empty = product_json(line_items=[
            {"quantity": 5, "unit_price": 2, "amount": 10}])
        rec = Recorder([envelope(empty)] * 6)
        monkeypatch.setattr(openrouter_client, "_chat_completion", rec)

        result = process_file(free_goods_pdf, pf_cfg(), logger)

        assert result.error is True               # the one hard failure
        assert "no usable product rows" in result.review_reason
        assert "provider_failure" not in safe_review_categories(result)


# --- semantic safety is unchanged ---------------------------------------------

class TestSemanticSafetyPreserved:
    def _shifted(self):
        """Quantity carrying the printed line total (M10 shift)."""
        return product_json(line_items=[
            {"item_code": c, "description": d, "quantity": a,
             "unit_price": u, "amount": a * u}
            for c, d, q, u, a in FREE_GOODS_ROWS])

    def test_column_shift_still_rejects_attempt(self, logger, free_goods_pdf,
                                                monkeypatch):
        rec = Recorder([envelope(self._shifted()),
                        envelope(product_json())])
        monkeypatch.setattr(openrouter_client, "_chat_completion", rec)

        result = process_file(free_goods_pdf, pf_cfg(), logger)

        assert len(text_calls(rec)) == 2          # escalated past the bad one
        assert result.error is False
        assert sum(i.quantity for i in result.invoice.line_items) == FG_QTY_TOTAL
        rejected = [r for r in result.usage_records if not r.accepted]
        assert [r.rejection_category for r in rejected] == \
            ["line_item_semantic_mismatch"]

    def test_missing_value_is_not_corruption(self, logger, free_goods_pdf,
                                             monkeypatch):
        """A missing quantity warns; a quantity holding the line total is
        still rejected. The two must not be conflated."""
        rows = [dict(zip(("item_code", "description", "quantity",
                          "unit_price", "amount"), r))
                for r in FREE_GOODS_ROWS]
        for row in rows:
            row["quantity"] = None
        rec = Recorder([envelope(product_json(line_items=rows))])
        monkeypatch.setattr(openrouter_client, "_chat_completion", rec)

        result = process_file(free_goods_pdf, pf_cfg(), logger)

        assert len(rec.calls) == 1                # accepted with warnings
        assert result.error is False
        assert all("quantity" in v for v in result.line_missing.values())

    def test_provider_failure_still_classified_as_such(self, logger,
                                                       free_goods_pdf,
                                                       monkeypatch):
        from invoice_extractor.provider import ProviderError

        def boom(*a, **k):
            raise ProviderError("upstream unavailable (HTTP 503)",
                                category="server_error", http_status=503)
        monkeypatch.setattr(openrouter_client, "_chat_completion", boom)

        result = process_file(free_goods_pdf, pf_cfg(1,
                                                     openrouter_vision_models=()),
                              logger)

        assert result.error is True
        assert "provider_failure" in safe_review_categories(result)

    def test_vision_fallback_still_bounded_to_one(self, logger,
                                                  free_goods_pdf,
                                                  monkeypatch):
        """No usable rows from any text model -> exactly one vision try."""
        empty = product_json(line_items=[{"quantity": 1, "unit_price": 2,
                                          "amount": 2}])
        rec = Recorder([envelope(empty), envelope(empty),
                        envelope(product_json())])
        monkeypatch.setattr(openrouter_client, "_chat_completion", rec)

        result = process_file(free_goods_pdf, pf_cfg(), logger)

        assert len(text_calls(rec)) == 2
        assert len(vision_calls(rec)) == 1        # exactly one
        assert result.vision_fallback_used is True
        assert result.error is False

    def test_attempt_and_cost_accounting(self, logger, free_goods_pdf,
                                         monkeypatch):
        rec = Recorder([envelope(self._shifted(), cost=0.001),
                        envelope(product_json(), cost=0.002)])
        monkeypatch.setattr(openrouter_client, "_chat_completion", rec)

        result = process_file(free_goods_pdf, pf_cfg(), logger)

        assert [float(r.cost_usd) for r in result.usage_records] == \
            [0.001, 0.002]
        assert [r.accepted for r in result.usage_records] == [False, True]


# --- workbook -----------------------------------------------------------------

class TestWorkbook:
    def _result(self, logger, pdf, monkeypatch, payload=None):
        rec = Recorder([envelope(payload or product_json())])
        monkeypatch.setattr(openrouter_client, "_chat_completion", rec)
        return process_file(pdf, pf_cfg(), logger)

    def test_columns_are_backward_compatible(self):
        """Pre-M11 columns keep their positions; new ones are appended."""
        assert INVOICE_COLUMNS[-3:] == ["document_type", "missing_fields",
                                        "validation_warnings"]
        assert LINE_ITEM_COLUMNS[-2:] == ["missing_fields",
                                          "validation_warnings"]
        assert LINE_ITEM_COLUMNS[:10] == [
            "invoice_id", "line_number", "source_file", "line_no",
            "item_code", "barcode", "description", "quantity", "unit_price",
            "amount"]

    def test_missing_cells_blank_highlighted_numbers_stay_numeric(
            self, logger, free_goods_pdf, monkeypatch, tmp_path):
        from openpyxl import load_workbook

        rows = [dict(zip(("item_code", "description", "quantity",
                          "unit_price", "amount"), r))
                for r in FREE_GOODS_ROWS]
        rows[0]["quantity"] = None
        result = self._result(logger, free_goods_pdf, monkeypatch,
                              product_json(line_items=rows))
        out = tmp_path / "wb.xlsx"
        export_workbook([result], out)

        book = load_workbook(out)
        inv = book["Invoices"]
        headers = [c.value for c in inv[1]]
        seller_col = headers.index("seller_name") + 1
        cell = inv.cell(row=2, column=seller_col)
        assert cell.value is None                          # blank, not "n/a"
        assert cell.fill.start_color.rgb == "FFFDF3C7"     # highlighted
        assert inv.cell(row=2, column=headers.index("document_type") + 1
                        ).value == "free_goods_support_list"

        li = book["LineItems"]
        li_headers = [c.value for c in li[1]]
        qty_col = li_headers.index("quantity") + 1
        assert li.cell(row=2, column=qty_col).value is None
        assert li.cell(row=2, column=qty_col).fill.start_color.rgb == \
            "FFFDF3C7"
        # untouched numeric cells stay numeric (never coerced to text)
        assert isinstance(li.cell(row=3, column=qty_col).value, (int, float))
        amount_col = li_headers.index("amount") + 1
        assert isinstance(li.cell(row=2, column=amount_col).value,
                          (int, float))

    def test_needs_review_populated_while_line_items_kept(
            self, logger, free_goods_pdf, monkeypatch, tmp_path):
        from openpyxl import load_workbook

        result = self._result(logger, free_goods_pdf, monkeypatch)
        out = tmp_path / "wb2.xlsx"
        export_workbook([result], out)

        book = load_workbook(out)
        review = list(book["NeedsReview"].iter_rows(values_only=True))
        assert len(review) == 2                    # header + one entry
        assert "missing document metadata" in review[1][4]
        lines = list(book["LineItems"].iter_rows(values_only=True))
        assert len(lines) == 4                     # header + 3 product rows

    def test_free_of_charge_total_preserved_not_zeroed(self, logger,
                                                       free_goods_pdf,
                                                       monkeypatch):
        result = self._result(logger, free_goods_pdf, monkeypatch)
        assert result.invoice.total_amount == FG_AMOUNT_TOTAL   # as printed
        assert result.free_of_charge is True
        assert "free-of-charge" in result.review_reason
        assert "free_of_charge_value_ambiguity" in \
            safe_review_categories(result)


class TestUiCounters:
    def test_counters_show_extracted_and_review_not_failed(self, logger,
                                                           free_goods_pdf,
                                                           monkeypatch):
        rec = Recorder([envelope(product_json())])
        monkeypatch.setattr(openrouter_client, "_chat_completion", rec)
        result = process_file(free_goods_pdf, pf_cfg(), logger)

        extracted = 0 if result.error else 1
        failed = 1 if result.error else 0
        assert (extracted, int(result.needs_review), failed) == (1, 1, 0)


class TestPromptContract:
    def test_prompt_forbids_sold_to_as_seller(self):
        """A document with no issuing party must leave seller_name null
        rather than promoting the sold-to party (live models did exactly
        that on a free-goods list that names only a SOLD-TO party)."""
        from invoice_extractor.prompts import RULES
        assert "never fall back to the sold-to" in RULES.lower()
        assert "is the buyer, never the seller" in RULES.lower()

    def test_prompt_prioritizes_line_items_for_product_documents(self):
        from invoice_extractor.prompts import RULES
        low = RULES.lower()
        assert "packing list" in low and "delivery note" in low
        assert "line items" in low


class TestNoHardcoding:
    def test_sources_carry_no_document_specific_values(self):
        root = Path(__file__).resolve().parent.parent
        for rel in ("invoice_extractor/schema.py",
                    "invoice_extractor/pipeline.py",
                    "invoice_extractor/excel_export.py",
                    "invoice_extractor/prompts.py"):
            src = (root / rel).read_text(encoding="utf-8")
            for banned in ("4006678", "IMAGINEX", "26177628", "44500",
                           "16862503", "NFRS"):
                assert banned not in src, f"{rel}: {banned}"
