"""Organization catalog (Build 11): the approved name-to-ID mapping, its
integrity, and the reference workbook cross-check. Offline only."""

from pathlib import Path

import pytest

from apps.web.transfer import organizations as orgs

ROOT = Path(__file__).resolve().parent.parent

APPROVED = {
    "Brooks Brothers Hong Kong": "100017",
    "Brooks Brothers Macao": "100018",
    "Brooks Brothers Malaysia": "100023",
    "Brooks Brothers Singapore": "100022",
    "Brooks Brothers Taiwan": "100021",
    "IMAGINEX Hong Kong": "100009",
    "IMAGINEX Macao": "100010",
    "IMAGINEX Singapore": "100014",
    "IMAGINEX Taiwan": "100012",
    "Joyce Beauty": "100007",
    "Sacai Hong Kong": "100024",
}


class TestCatalog:
    def test_all_eleven_approved_organizations(self):
        assert len(orgs.ORGANIZATIONS) == 11
        assert {o.name: o.org_id for o in orgs.ORGANIZATIONS} == APPROVED

    def test_names_and_ids_unique(self):
        names = [o.name for o in orgs.ORGANIZATIONS]
        ids = [o.org_id for o in orgs.ORGANIZATIONS]
        assert len(set(names)) == len(names)
        assert len(set(ids)) == len(ids)

    def test_ids_are_strings_matching_the_orgid_dto_type(self):
        # the tracked spec types orgId as string; the catalog matches
        assert all(isinstance(o.org_id, str) and o.org_id.isdigit()
                   for o in orgs.ORGANIZATIONS)

    def test_lookup_helpers(self):
        assert orgs.by_name("IMAGINEX Hong Kong").org_id == "100009"
        assert orgs.by_id("100012").name == "IMAGINEX Taiwan"
        assert orgs.by_id(100007).name == "Joyce Beauty"   # int tolerated
        assert orgs.by_name("Nope") is None
        assert orgs.is_valid_org_id("999999") is False
        assert orgs.is_valid_org_id(None) is False

    def test_correct_spelling_everywhere(self):
        source = (ROOT / "apps" / "web" / "transfer"
                  / "organizations.py").read_text(encoding="utf-8")
        assert "Organization ID" in source
        assert "Orginazation" not in source        # workbook header typo

    def test_no_runtime_excel_dependency(self):
        # the catalog is typed, deterministic code - it never reads the
        # reference workbook (the docstring may MENTION it)
        source = (ROOT / "apps" / "web" / "transfer"
                  / "organizations.py").read_text(encoding="utf-8")
        for banned in ("import openpyxl", "load_workbook", "import pandas",
                       "read_excel", "open("):
            assert banned not in source, banned
        import apps.web.transfer.organizations as module
        assert not hasattr(module, "openpyxl")


@pytest.mark.skipif(
    not (ROOT / "organization.xlsx").exists(),
    reason="reference workbook not present on this machine")
class TestReferenceWorkbook:
    def test_workbook_matches_the_approved_catalog(self):
        from openpyxl import load_workbook
        book = load_workbook(ROOT / "organization.xlsx", read_only=True)
        rows = list(book[book.sheetnames[0]].iter_rows(values_only=True))
        mapping = {str(name).strip(): str(org_id).strip()
                   for name, org_id in rows[1:] if name and org_id}
        assert mapping == APPROVED
