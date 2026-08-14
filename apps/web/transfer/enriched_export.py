"""Final enriched product lines: shared column schema + Excel export
(Build 10).

One row per delivery-note line, straight from the persisted enrichment
artifact. The schema is FIXED and identical across jobs (empty Analysis
Code / Composition columns are kept, never dropped), and it is shared by
the UI table and the Excel export so the two can never drift. Friendly
display names only - the gateway's misspelled wire names (compositon1)
never reach the user.

The export is purely local and read-only: no API call, no job-state
change, no checkpoint touch, no temporary files (bytes in memory only).
Identifiers are written as text (leading zeros preserved, no scientific
notation); quantities are whole numbers; prices are numeric.
"""

import re
from io import BytesIO

from apps.web.job_manager import JobError

WORKSHEET_NAME = "Enriched Product Lines"

# (header, kind) - kind: "text" (identifier, format '@'), "int", "price",
# "plain" (free text). Stable order: identity/issues left, source, API,
# prices BEFORE the attribute block, then AC01..AC15 and Composition 1..4.
EXPORT_COLUMNS: tuple[tuple[str, str], ...] = tuple(
    [("Line", "text"), ("Carton", "text"), ("Status", "plain"),
     ("Via", "plain"), ("Issues", "int"),
     ("Source EAN", "text"), ("Source item", "text"),
     ("Source description", "plain"), ("Source quantity", "int"),
     ("API EAN", "text"), ("API item", "text"), ("API color", "text"),
     ("API size", "text"), ("API description", "plain"),
     ("Original price", "price"), ("Discount price", "price")]
    + [(f"AC{i:02d}", "plain") for i in range(1, 16)]
    + [(f"Composition {i}", "plain") for i in range(1, 5)])

DEFAULT_VISIBLE_COLUMNS = ("Line", "Carton", "Status", "Via", "Source EAN",
                           "API item", "API color", "API size",
                           "API description", "Original price",
                           "Discount price", "Issues")

ATTRIBUTE_COLUMNS = tuple(f"AC{i:02d}" for i in range(1, 16)) + tuple(
    f"Composition {i}" for i in range(1, 5))


def export_filename(job_id: str) -> str:
    safe = re.sub(r"[^A-Za-z0-9_-]", "-", job_id)
    return f"Transfer_{safe}_Enriched_Product_Lines.xlsx"


def _price(value):
    try:
        return float(value) if value not in (None, "") else None
    except (TypeError, ValueError):
        return None


def _whole(value):
    try:
        return int(value) if value not in (None, "") else None
    except (TypeError, ValueError):
        return None


def build_rows(enrichment: dict) -> list[dict]:
    """One dict per delivery-note line, in persisted order, keyed by the
    EXPORT_COLUMNS headers. Identifier values stay strings."""
    products = enrichment.get("products", [])
    rows = []
    for line in enrichment.get("line_enrichments", []):
        product = (products[line["product_ref"]]
                   if line.get("product_ref") is not None else {})
        source = line.get("source", {})
        row = {
            "Line": line.get("line_id"),
            "Carton": line.get("original_carton_number"),
            "Status": line.get("status"),
            "Via": line.get("matched_via") or "-",
            "Issues": int(line.get("comparison_issue_count") or 0),
            "Source EAN": source.get("ean"),
            "Source item": source.get("item_code"),
            "Source description": source.get("description"),
            "Source quantity": _whole(source.get("quantity")),
            "API EAN": product.get("ean"),
            "API item": product.get("item_code"),
            "API color": product.get("color_code"),
            "API size": product.get("size_code"),
            "API description": product.get("item_desc"),
            "Original price": _price(product.get("original_retail_price")),
            "Discount price": _price(product.get("discount_price")),
        }
        for i in range(1, 16):
            row[f"AC{i:02d}"] = product.get(f"analysis_code_{i:02d}") or ""
        for i in range(1, 5):
            row[f"Composition {i}"] = (product.get(f"composition_{i:02d}")
                                       or "")
        rows.append(row)
    return rows


_WIDTHS = {"Line": 16, "Carton": 8, "Status": 11, "Via": 12, "Issues": 8,
           "Source EAN": 16, "Source item": 17, "Source description": 30,
           "Source quantity": 9, "API EAN": 18, "API item": 17,
           "API color": 10, "API size": 9, "API description": 30,
           "Original price": 13, "Discount price": 13}


def build_enriched_workbook_bytes(job_id: str) -> bytes:
    """Build the .xlsx entirely in memory from the persisted enrichment.
    Read-only: no API call, no job-state change, no checkpoint touch, no
    files written anywhere."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    from apps.web.transfer import product_lookup as pl

    enrichment = pl.load_enrichment(job_id)
    if enrichment is None:
        raise JobError("No product lookup result exists for this job.")
    rows = build_rows(enrichment)

    book = Workbook()
    sheet = book.active
    sheet.title = WORKSHEET_NAME
    bold = Font(bold=True)
    for col, (header, _) in enumerate(EXPORT_COLUMNS, start=1):
        cell = sheet.cell(row=1, column=col, value=header)
        cell.font = bold
        width = _WIDTHS.get(header, 10 if header.startswith("AC") else 16)
        sheet.column_dimensions[get_column_letter(col)].width = width
    for row_index, row in enumerate(rows, start=2):
        for col, (header, kind) in enumerate(EXPORT_COLUMNS, start=1):
            value = row.get(header)
            if value in (None, ""):
                continue                          # blanks stay blank cells
            cell = sheet.cell(row=row_index, column=col)
            if kind == "text":
                cell.value = str(value)
                cell.number_format = "@"          # leading zeros preserved
            elif kind == "int":
                cell.value = int(value)
            elif kind == "price":
                cell.value = float(value)
            else:
                cell.value = str(value)
    sheet.freeze_panes = "A2"
    last_col = get_column_letter(len(EXPORT_COLUMNS))
    sheet.auto_filter.ref = f"A1:{last_col}{max(len(rows) + 1, 1)}"

    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()
