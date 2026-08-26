"""Excel workbook export: Invoices / LineItems / NeedsReview sheets.

Decimal values are converted to float at this boundary only (Excel cells are
IEEE floats); Decimal is used everywhere upstream.
"""

from decimal import Decimal
from pathlib import Path

import pandas as pd
from openpyxl.styles import PatternFill

from invoice_extractor.pdf_utils import format_page_ranges
from invoice_extractor.pipeline import InvoiceResult
from invoice_extractor.schema import HEADER_FIELDS, LINE_ITEM_FIELDS, suspicious_line_item_rows

# Page columns use the documented human-readable range format ("1-2,5-7"),
# never Python list syntax.
PROVENANCE_COLUMNS = [
    "source_file",
    "page_count",
    "document_classification",
    "extraction_method",
    "provider",
    "model",
    "text_pages",
    "image_pages",
    "blank_pages",
    "failed_pages",
    "vision_chunk_count",
]

# M11: document_type/missing_fields/validation_warnings are APPENDED after the
# existing columns so every pre-M11 column keeps its position and meaning.
INVOICE_COLUMNS = (
    ["invoice_id"] + HEADER_FIELDS + PROVENANCE_COLUMNS
    + ["needs_review", "review_reason", "document_type", "missing_fields",
       "validation_warnings"]
)

LINE_ITEM_COLUMNS = (["invoice_id", "line_number", "source_file"]
                     + LINE_ITEM_FIELDS
                     + ["missing_fields", "validation_warnings"])

# A focused, reviewer-facing sheet - deliberately NOT a slice of INVOICE_COLUMNS.
# line_numbers/line_descriptions are None unless suspicious_line_item_rows()
# found specific rows to point at (e.g. the hallucinated-header-row guardrail);
# other review reasons (missing fields, totals inconclusive, ...) are invoice-
# level and leave those two columns blank.
NEEDS_REVIEW_COLUMNS = [
    "invoice_id", "source_file", "invoice_number", "needs_review",
    "review_reason", "line_numbers", "line_descriptions",
]


def _cell(value):
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, list):
        return format_page_ranges(value)
    return value


# Restrained warning fill: pale yellow, readable in print and on screen.
_MISSING_FILL = PatternFill("solid", start_color="FFFDF3C7",
                            end_color="FFFDF3C7")


def _highlight_missing(ws, df, results, *, level: str) -> None:
    """Shade blank cells whose value the source document never printed.

    Only cells named by the per-document/per-row missing-field lists are
    shaded - a cell that is blank for any other reason is left alone, so the
    fill always means exactly one thing to a reviewer."""
    columns = {name: idx for idx, name in enumerate(df.columns, start=1)}
    if level == "document":
        for row_idx, res in enumerate(results, start=2):   # row 1 = header
            for field in getattr(res, "document_missing", []) or []:
                col = columns.get(field)
                if col:
                    ws.cell(row=row_idx, column=col).fill = _MISSING_FILL
        return
    excel_row = 1                                          # row 1 = header
    for res in results:
        line_missing = getattr(res, "line_missing", {}) or {}
        for line_no in range(1, len(res.invoice.line_items) + 1):
            excel_row += 1
            for field in line_missing.get(line_no, []):
                col = columns.get(field)
                if col:
                    ws.cell(row=excel_row, column=col).fill = _MISSING_FILL


def _document_warnings(res) -> str | None:
    """Reviewer-facing warning summary for one document (M11). Warnings are
    NOT failures: the product rows are present and exported."""
    parts = []
    missing = getattr(res, "document_missing", []) or []
    if missing:
        parts.append("document metadata not printed: " + ", ".join(missing))
    line_missing = getattr(res, "line_missing", {}) or {}
    if line_missing:
        rows = ", ".join(str(r) for r in sorted(line_missing)[:12])
        parts.append(f"line-item fields missing on row(s) {rows}")
    if getattr(res, "free_of_charge", False):
        parts.append("declared free-of-charge: printed values may be "
                     "declared/reference amounts, payable amount unclear")
    return " | ".join(parts) or None


def export_workbook(results: list[InvoiceResult], output_path: str | Path) -> Path:
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    invoice_rows = []
    line_item_rows = []
    needs_review_rows = []

    for i, res in enumerate(results, start=1):
        invoice_id = f"INV-{i:04d}"
        row = {"invoice_id": invoice_id}
        row.update({f: _cell(getattr(res.invoice, f)) for f in HEADER_FIELDS})
        row.update(
            source_file=res.source_file,
            page_count=res.page_count,
            document_classification=res.document_classification,
            extraction_method=res.extraction_method,
            provider=res.provider,
            model=res.model,
            text_pages=_cell(res.text_pages),
            image_pages=_cell(res.image_pages),
            blank_pages=_cell(res.blank_pages),
            failed_pages=_cell(res.failed_pages),
            vision_chunk_count=res.vision_chunk_count,
            needs_review=res.needs_review,
            review_reason=res.review_reason,
            document_type=getattr(res, "document_type", None),
            missing_fields=", ".join(getattr(res, "document_missing", []))
            or None,
            validation_warnings=_document_warnings(res),
        )
        invoice_rows.append(row)

        line_missing = getattr(res, "line_missing", {}) or {}
        for j, item in enumerate(res.invoice.line_items, start=1):
            li = {"invoice_id": invoice_id, "line_number": j, "source_file": res.source_file}
            li.update({f: _cell(getattr(item, f)) for f in LINE_ITEM_FIELDS})
            missing = line_missing.get(j) or []
            li["missing_fields"] = ", ".join(missing) or None
            li["validation_warnings"] = (
                "not printed on the document: " + ", ".join(missing)
                if missing else None)
            line_item_rows.append(li)

        if res.needs_review:
            suspicious = suspicious_line_item_rows(res.invoice.line_items)
            needs_review_rows.append({
                "invoice_id": invoice_id,
                "source_file": res.source_file,
                "invoice_number": row["invoice_number"],
                "needs_review": res.needs_review,
                "review_reason": res.review_reason,
                "line_numbers": ", ".join(str(n) for n in suspicious) or None,
                "line_descriptions": "; ".join(
                    res.invoice.line_items[n - 1].description or "(no description)"
                    for n in suspicious
                ) or None,
            })

    invoices_df = pd.DataFrame(invoice_rows, columns=INVOICE_COLUMNS)
    line_items_df = pd.DataFrame(line_item_rows, columns=LINE_ITEM_COLUMNS)
    needs_review_df = pd.DataFrame(needs_review_rows, columns=NEEDS_REVIEW_COLUMNS)

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        invoices_df.to_excel(writer, sheet_name="Invoices", index=False)
        line_items_df.to_excel(writer, sheet_name="LineItems", index=False)
        needs_review_df.to_excel(writer, sheet_name="NeedsReview", index=False)

        # M11: highlight cells the DOCUMENT did not supply. Cells are left
        # genuinely blank (never "n/a" text) and typed values elsewhere are
        # untouched, so numeric columns stay numeric and sortable - only the
        # fill changes.
        _highlight_missing(writer.sheets["Invoices"], invoices_df,
                           results, level="document")
        _highlight_missing(writer.sheets["LineItems"], line_items_df,
                           results, level="line")

        # Reasonable column widths for quick triage; tolerate null/NaN values.
        for sheet_name, df in (
            ("Invoices", invoices_df),
            ("LineItems", line_items_df),
            ("NeedsReview", needs_review_df),
        ):
            ws = writer.sheets[sheet_name]
            for col_idx, col_name in enumerate(df.columns, start=1):
                lengths = [
                    len(str(v)) for v in df[col_name].head(200) if not pd.isna(v)
                ]
                max_len = max([len(str(col_name))] + lengths)
                ws.column_dimensions[
                    ws.cell(row=1, column=col_idx).column_letter
                ].width = min(max_len + 2, 60)

    return output_path
